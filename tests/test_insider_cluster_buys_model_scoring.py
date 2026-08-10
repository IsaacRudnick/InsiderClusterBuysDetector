"""Tests for insider_cluster_buys.py's model-based live scoring wiring
(Part 1 + Part 2 of the conviction_score -> trained-model migration):

  - _resolve_latest_artifact / _price_fetch_enabled: the small pure-ish
    helpers around opt-out and artifact discovery.
  - build_live_price_universe: graceful degradation when there is nothing
    to fetch or the fetch fails, via a FakePriceUniverse double (no network).
  - score_clusters_with_model / attach_model_scores: the orchestration that
    scores every cluster and reports feature-availability honestly, with a
    tiny REAL ProductionBundle (real fitted LGBMClassifier, same pattern as
    tests/test_live_score.py's make_bundle) so the predict_proba path is
    exercised for real -- cheaply, no network, no research_data/ dependency
    (LIVE_SCORE_MODEL_PATH / LIVE_SCORE_HISTORY_PATH env overrides point at
    tmp_path files instead).
  - write_excel / build_payload: the conviction_score columns are gone from
    the Flagged Clusters sheet, replaced by the model's percentile/verdict/
    coverage/reasons, and model_info flows into the payload.

Also guards the task's explicit scope limit: DEFAULT_WEIGHTS, _score_cluster,
_component_flags, and _build_cluster's own conviction_score/_label/
_contributions fields must still exist unchanged -- backtest/state.py
depends on them as the experimental control.
"""

from __future__ import annotations

import os
from datetime import date, timedelta

import numpy as np
import openpyxl
import pandas as pd
import pytest

import insider_cluster_buys as ics
from research import model as rm


# ---------------------------------------------------------------------------
# Fixtures / builders
# ---------------------------------------------------------------------------
def _qual_row(
    *, ticker: str = "ABC", issuer_cik: str = "CIK1", issuer_name: str = "Issuer Co",
    owner_cik: str = "O1", owner_name: str = "Owner One", owner_roles: str = "Director",
    is_director: bool = True, is_officer: bool = False, is_ten_pct: bool = False,
    transaction_date: str, filing_date: str, shares: float = 1000.0,
    price_per_share: float = 10.0, value: float = 10000.0,
    pct_of_prior_stake=None, footnote_text: str = "", is_10b5_1: bool = False,
) -> dict:
    """One row shaped exactly like insider_cluster_buys.extract_qualifying_rows'
    output -- the real input _build_cluster/detect_clusters consume."""
    return {
        "adsh": "0000000000-26-000001", "form_type": "4", "filing_url": "https://example.com/filing",
        "issuer_cik": issuer_cik, "issuer_name": issuer_name, "ticker": ticker,
        "owner_cik": owner_cik, "owner_name": owner_name, "owner_roles": owner_roles,
        "is_director": is_director, "is_officer": is_officer, "is_ten_percent_owner": is_ten_pct,
        "transaction_date": transaction_date, "transaction_code": "P", "acquired_disposed": "A",
        "shares": shares, "price_per_share": price_per_share, "value": value,
        "shares_owned_after": (shares or 0) * 10, "pct_of_prior_stake": pct_of_prior_stake,
        "filing_date": filing_date, "footnote_text": footnote_text, "is_10b5_1": is_10b5_1,
    }


def make_cluster(ticker: str = "ABC", n_owners: int = 3) -> dict:
    window = [
        _qual_row(
            ticker=ticker, owner_cik=f"O{i}", owner_name=f"Owner {i}",
            transaction_date=f"2026-08-{i + 1:02d}", filing_date=f"2026080{i + 2}",
            price_per_share=10.0 + i, shares=1000.0 + 100 * i,
            value=(10.0 + i) * (1000.0 + 100 * i), pct_of_prior_stake=5.0 + i,
            is_director=(i % 2 == 0), is_officer=(i % 2 == 1),
        )
        for i in range(n_owners)
    ]
    return ics._build_cluster(window)


def make_bundle(feature_cols: list[str], *, seed: int = 0, n: int = 300) -> rm.ProductionBundle:
    """Tiny real ProductionBundle -- same construction as
    tests/test_live_score.py's own make_bundle, duplicated here per this
    project's own test-file convention (each file owns its fixtures)."""
    rng = np.random.default_rng(seed)
    df = pd.DataFrame({c: rng.standard_normal(n) for c in feature_cols})
    df["entry_idx"] = np.arange(n)
    df[rm.LABEL_COL] = rng.standard_normal(n) * 0.05
    df.loc[: n // 5, rm.LABEL_COL] = rng.uniform(0.25, 0.6, n // 5 + 1)

    fast_params = dict(
        n_estimators=40, num_leaves=7, min_child_samples=8, learning_rate=0.1,
        subsample=0.9, subsample_freq=1, colsample_bytree=0.9, reg_lambda=1.0,
        random_state=0, verbosity=-1,
    )
    target = (df[rm.LABEL_COL] > rm.TAIL_THRESH).astype(int)
    model = rm._fit_classifier(df, feature_cols, target, fast_params)
    scores = np.sort(rm._predict_proba_positive(model, df[feature_cols], fallback_rate=float("nan")))
    return rm.ProductionBundle(
        model=model, feature_cols=list(feature_cols), training_scores=scores,
        provenance={"source_dataset_path": "synthetic", "n_rows": n}, config={},
    )


def make_issuer_history_frame() -> pd.DataFrame:
    """Minimal frame shaped exactly like what
    research.live_score._issuer_ticker_history_features reads: ticker,
    issuer_cik, event_day, and the label column. Deliberately NOT run
    through research.model.load_research_dataset's full schema validation
    (that validation is test_model.py's job) -- attach_model_scores' own
    unit tests below monkeypatch rm.load_research_dataset to hand this
    back directly, so only the wiring is under test here."""
    return pd.DataFrame({
        "ticker": ["ABC", "ABC", "XYZ"],
        "issuer_cik": ["CIK1", "CIK1", "CIK9"],
        "event_day": [date(2024, 1, 1), date(2025, 1, 1), date(2024, 6, 1)],
        rm.LABEL_COL: [0.05, 0.10, -0.02],
    })


class FakePriceUniverse:
    """Minimal PriceUniverse double for build_live_price_universe tests --
    no disk, no network. Mirrors tests/test_live_score.py's own double."""

    def __init__(self, *, raise_on_ensure: bool = False):
        self.frames: dict[str, object] = {}
        self._raise_on_ensure = raise_on_ensure
        self.ensure_calls: list[tuple] = []
        self.finalized = False

    def ensure(self, tickers, start, end) -> None:
        self.ensure_calls.append((tuple(tickers), start, end))
        if self._raise_on_ensure:
            raise RuntimeError("simulated fetch failure")
        for t in tickers:
            self.frames[t] = object()

    def finalize(self) -> None:
        self.finalized = True


# ---------------------------------------------------------------------------
# Scope-limit guard: DEFAULT_WEIGHTS / _score_cluster / _component_flags /
# _build_cluster's conviction_* fields must be untouched.
# ---------------------------------------------------------------------------
class TestScopeLimitUntouched:
    def test_default_weights_and_scorer_functions_still_exist(self):
        assert isinstance(ics.DEFAULT_WEIGHTS, dict) and ics.DEFAULT_WEIGHTS
        assert callable(ics._score_cluster)
        assert callable(ics._component_flags)

    def test_build_cluster_still_populates_conviction_fields(self):
        c = make_cluster()
        assert "conviction_score" in c
        assert "conviction_label" in c
        assert "conviction_contributions" in c
        assert c["conviction_label"] in ("conviction", "mixed", "routine")


# ---------------------------------------------------------------------------
# _resolve_latest_artifact
# ---------------------------------------------------------------------------
class TestResolveLatestArtifact:
    def test_returns_none_when_nothing_matches(self, tmp_path):
        assert ics._resolve_latest_artifact("production_model_*.joblib", out_dir=str(tmp_path)) is None

    def test_picks_the_newest_by_mtime(self, tmp_path):
        older = tmp_path / "production_model_100rows_20260101.joblib"
        newer = tmp_path / "production_model_200rows_20260202.joblib"
        older.write_bytes(b"old")
        newer.write_bytes(b"new")
        os.utime(older, (1000, 1000))
        os.utime(newer, (2000, 2000))
        assert ics._resolve_latest_artifact("production_model_*.joblib", out_dir=str(tmp_path)) == str(newer)


# ---------------------------------------------------------------------------
# _price_fetch_enabled
# ---------------------------------------------------------------------------
class TestPriceFetchEnabled:
    def test_enabled_by_default(self, monkeypatch):
        monkeypatch.delenv(ics.LIVE_SCORE_FETCH_PRICES_ENV, raising=False)
        monkeypatch.setattr("sys.argv", ["insider_cluster_buys.py"])
        assert ics._price_fetch_enabled() is True

    def test_disabled_by_env_var(self, monkeypatch):
        monkeypatch.setenv(ics.LIVE_SCORE_FETCH_PRICES_ENV, "0")
        monkeypatch.setattr("sys.argv", ["insider_cluster_buys.py"])
        assert ics._price_fetch_enabled() is False

    def test_disabled_by_false_string(self, monkeypatch):
        monkeypatch.setenv(ics.LIVE_SCORE_FETCH_PRICES_ENV, "false")
        monkeypatch.setattr("sys.argv", ["insider_cluster_buys.py"])
        assert ics._price_fetch_enabled() is False

    def test_disabled_by_cli_flag(self, monkeypatch):
        monkeypatch.delenv(ics.LIVE_SCORE_FETCH_PRICES_ENV, raising=False)
        monkeypatch.setattr("sys.argv", ["insider_cluster_buys.py", "--no-prices"])
        assert ics._price_fetch_enabled() is False


# ---------------------------------------------------------------------------
# build_live_price_universe -- graceful degradation, no network
# ---------------------------------------------------------------------------
class TestBuildLivePriceUniverse:
    def test_no_tickers_returns_none_without_constructing_anything(self, monkeypatch):
        def _boom(*a, **k):
            raise AssertionError("PriceUniverse should not be constructed with no tickers")
        monkeypatch.setattr("backtest.prices.PriceUniverse", _boom)
        assert ics.build_live_price_universe([{"ticker": "", "cluster_end": "2026-08-01"}]) is None

    def test_successful_fetch_returns_populated_universe(self, monkeypatch):
        fake = FakePriceUniverse()
        monkeypatch.setattr("backtest.prices.PriceUniverse", lambda: fake)
        clusters = [{"ticker": "ABC", "cluster_end": "2026-08-01"}]
        result = ics.build_live_price_universe(clusters)
        assert result is fake
        assert fake.finalized is True
        assert "ABC" in fake.frames

    def test_failed_fetch_degrades_to_none(self, monkeypatch):
        fake = FakePriceUniverse(raise_on_ensure=True)
        monkeypatch.setattr("backtest.prices.PriceUniverse", lambda: fake)
        clusters = [{"ticker": "ABC", "cluster_end": "2026-08-01"}]
        assert ics.build_live_price_universe(clusters) is None


# ---------------------------------------------------------------------------
# score_clusters_with_model
# ---------------------------------------------------------------------------
class TestScoreClustersWithModel:
    def test_no_bundle_sets_every_model_score_to_none(self):
        clusters = [make_cluster("AAA"), make_cluster("BBB")]
        ics.score_clusters_with_model(clusters, bundle=None)
        assert all(c["model_score"] is None for c in clusters)

    def test_real_bundle_populates_model_score_shape(self):
        bundle = make_bundle(list(rm.FEATURE_COLS))
        clusters = [make_cluster("AAA")]
        ics.score_clusters_with_model(clusters, bundle=bundle)
        ms = clusters[0]["model_score"]
        assert ms is not None
        assert set(ms) == {
            "raw_score", "percentile", "verdict", "n_features_available",
            "n_features_total", "missing_features", "n_training_scores", "factors",
        }
        assert ms["verdict"] in ("top_decile", "no_edge", "unavailable")
        assert ms["n_features_total"] == len(rm.FEATURE_COLS)
        assert ms["n_training_scores"] == len(bundle.training_scores)
        assert len(ms["factors"]) == 6  # REASON_PANEL_FACTORS length
        # No PriceUniverse / issuer_history passed -> only the always-
        # available 31 features should be computed.
        from research import live_score as ls
        assert ms["n_features_available"] == len(ls.ALWAYS_AVAILABLE_FEATURE_COLS)

    def test_cluster_with_no_transactions_is_not_scored(self):
        bundle = make_bundle(list(rm.FEATURE_COLS))
        c = make_cluster("AAA")
        c["transactions"] = []
        ics.score_clusters_with_model([c], bundle=bundle)
        assert c["model_score"] is None

    def test_scoring_failure_for_one_cluster_does_not_take_down_the_run(self, monkeypatch):
        bundle = make_bundle(list(rm.FEATURE_COLS))
        clusters = [make_cluster("AAA"), make_cluster("BBB")]

        import research.live_score as ls_mod
        real = ls_mod.score_live_cluster
        calls = {"n": 0}

        def _flaky(cluster, window, bundle, **kwargs):
            calls["n"] += 1
            if cluster.get("ticker") == "AAA":
                raise RuntimeError("simulated scoring failure")
            return real(cluster, window, bundle, **kwargs)

        monkeypatch.setattr(ls_mod, "score_live_cluster", _flaky)
        ics.score_clusters_with_model(clusters, bundle=bundle)
        assert clusters[0]["model_score"] is None
        assert clusters[1]["model_score"] is not None


# ---------------------------------------------------------------------------
# attach_model_scores -- full orchestration, hermetic via env-var overrides
# ---------------------------------------------------------------------------
class TestAttachModelScores:
    def test_no_bundle_on_disk_degrades_cleanly(self, tmp_path, monkeypatch):
        monkeypatch.setenv("LIVE_SCORE_MODEL_PATH", str(tmp_path / "nope.joblib"))
        monkeypatch.delenv("LIVE_SCORE_HISTORY_PATH", raising=False)
        clusters = [make_cluster("AAA")]
        info = ics.attach_model_scores(clusters)
        assert info["model_available"] is False
        assert info["n_training_scores"] == 0
        assert info["issuer_history_available"] is False
        assert clusters[0]["model_score"] is None

    def test_full_wiring_with_bundle_and_issuer_history_no_prices(self, tmp_path, monkeypatch):
        bundle = make_bundle(list(rm.FEATURE_COLS))
        bundle_path = tmp_path / "bundle.joblib"
        rm.save_production_bundle(bundle, str(bundle_path))

        history_path = tmp_path / "history.parquet"
        make_issuer_history_frame().to_parquet(history_path)

        monkeypatch.setenv("LIVE_SCORE_MODEL_PATH", str(bundle_path))
        monkeypatch.setenv("LIVE_SCORE_HISTORY_PATH", str(history_path))
        monkeypatch.setattr(rm, "load_research_dataset", lambda path: pd.read_parquet(path))
        monkeypatch.setenv(ics.LIVE_SCORE_FETCH_PRICES_ENV, "0")

        def _boom(*a, **k):
            raise AssertionError("price fetch must be skipped when disabled")
        monkeypatch.setattr(ics, "build_live_price_universe", _boom)

        clusters = [make_cluster("AAA")]
        info = ics.attach_model_scores(clusters)

        assert info["model_available"] is True
        assert info["n_training_scores"] == len(bundle.training_scores)
        assert info["issuer_history_available"] is True
        assert info["n_issuer_history_rows"] == 3
        assert info["price_fetch_enabled"] is False
        assert info["prices_loaded"] is False
        assert clusters[0]["model_score"] is not None
        # issuer_history was supplied -> issuer-history features on top of
        # the always-available set should now be computed.
        from research import live_score as ls
        expected = len(ls.ALWAYS_AVAILABLE_FEATURE_COLS) + len(ls.ISSUER_HISTORY_FEATURE_COLS)
        assert clusters[0]["model_score"]["n_features_available"] == expected

    def test_price_fetch_enabled_calls_build_live_price_universe(self, tmp_path, monkeypatch):
        bundle = make_bundle(list(rm.FEATURE_COLS))
        bundle_path = tmp_path / "bundle.joblib"
        rm.save_production_bundle(bundle, str(bundle_path))
        monkeypatch.setenv("LIVE_SCORE_MODEL_PATH", str(bundle_path))
        # Point at a path that does not exist rather than deleting the env
        # var: deleting it would fall back to _resolve_latest_artifact,
        # which would pick up this repo's own real research_data/ parquet
        # and make the test depend on what happens to be on disk.
        monkeypatch.setenv("LIVE_SCORE_HISTORY_PATH", str(tmp_path / "no_history_here.parquet"))
        monkeypatch.setenv(ics.LIVE_SCORE_FETCH_PRICES_ENV, "1")
        monkeypatch.setattr("sys.argv", ["insider_cluster_buys.py"])

        calls = {"n": 0}

        def _fake_universe(clusters):
            calls["n"] += 1
            return None

        monkeypatch.setattr(ics, "build_live_price_universe", _fake_universe)
        clusters = [make_cluster("AAA")]
        info = ics.attach_model_scores(clusters)
        assert calls["n"] == 1
        assert info["price_fetch_enabled"] is True
        assert info["prices_loaded"] is False


# ---------------------------------------------------------------------------
# write_excel / build_payload
# ---------------------------------------------------------------------------
class TestWriteExcelAndBuildPayload:
    def test_build_payload_includes_model_info(self):
        clusters = [make_cluster("AAA")]
        payload = ics.build_payload(clusters, (date(2026, 7, 1), date(2026, 8, 8)), model_info={"model_available": True})
        assert payload["model_info"] == {"model_available": True}

    def test_build_payload_defaults_model_info_to_empty_dict(self):
        payload = ics.build_payload([], (date(2026, 7, 1), date(2026, 8, 8)))
        assert payload["model_info"] == {}

    def test_write_excel_flagged_clusters_sheet_has_model_columns_not_conviction(self, tmp_path):
        bundle = make_bundle(list(rm.FEATURE_COLS))
        clusters = [make_cluster("AAA"), make_cluster("BBB")]
        ics.score_clusters_with_model(clusters, bundle=bundle)
        out_path = tmp_path / "out.xlsx"
        ics.write_excel(clusters, [], [], str(out_path), model_info={"model_available": True})
        assert out_path.exists()

        wb = openpyxl.load_workbook(str(out_path))
        ws = wb["Flagged Clusters"]
        header = [c.value for c in ws[1]]
        assert "Model Percentile" in header
        assert "Model Verdict" in header
        assert "Feature Coverage" in header
        assert "Model Reasons" in header
        assert "Signal Score" not in header
        assert "Signal Drivers" not in header
        assert "Model Coverage" in wb.sheetnames

    def test_write_excel_handles_unscored_clusters_without_crashing(self, tmp_path):
        clusters = [make_cluster("AAA")]
        ics.score_clusters_with_model(clusters, bundle=None)
        out_path = tmp_path / "out2.xlsx"
        ics.write_excel(clusters, [], [], str(out_path), model_info=None)
        assert out_path.exists()
